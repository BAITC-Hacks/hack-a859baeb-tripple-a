"""Bounded document readers retaining source coordinates for every text chunk."""
from __future__ import annotations

import hashlib
import io
import re
import unicodedata
import zipfile
from pathlib import PurePosixPath
from typing import Literal

from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pypdf import PdfReader

from .models import Chunk, Document

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_TEXT_CHARS = 500_000
MAX_PAGES = 300
MAX_ROWS = 30_000
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".xlsx"}


class IngestionError(ValueError):
    """A safe, user-facing document validation or extraction failure."""


def sanitize_filename(filename: str) -> str:
    """Keep a readable basename; never treat supplied names as filesystem paths."""
    name = PurePosixPath((filename or "document").replace("\\", "/")).name
    name = "".join(c for c in name if not unicodedata.category(c).startswith("C"))
    name = re.sub(r'[<>:"|?*]', "_", name).strip(" .")
    if not name:
        return "document"
    if len(name) > 180:
        suffix = PurePosixPath(name).suffix
        name = name[: 180 - len(suffix)] + suffix
    return name


def validate_file(filename: str, content: bytes) -> tuple[str, str]:
    name = sanitize_filename(filename)
    extension = PurePosixPath(name).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise IngestionError("Supported formats are PDF, DOCX and XLSX.")
    if not content:
        raise IngestionError(f"{name}: the file is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise IngestionError(f"{name}: exceeds the 20 MB per-file limit.")
    if extension == ".pdf":
        if not content[:1024].lstrip().startswith(b"%PDF-"):
            raise IngestionError(f"{name}: file contents do not match the PDF extension.")
    else:
        _validate_office_archive(name, extension, content)
    return name, extension[1:]


def _validate_office_archive(name: str, extension: str, content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 5000 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
                raise IngestionError(f"{name}: expanded Office archive exceeds the safety limit.")
            names = {e.filename for e in entries}
            expected = "word/document.xml" if extension == ".docx" else "xl/workbook.xml"
            if expected not in names or "[Content_Types].xml" not in names:
                raise IngestionError(f"{name}: contents do not match the {extension[1:].upper()} extension.")
            for entry in entries:
                path = PurePosixPath(entry.filename.replace("\\", "/"))
                if path.is_absolute() or ".." in path.parts:
                    raise IngestionError(f"{name}: archive contains an unsafe entry.")
                if entry.flag_bits & 1:
                    raise IngestionError(f"{name}: password-protected files are not supported.")
                if entry.file_size > 1_000_000 and entry.file_size / max(entry.compress_size, 1) > 500:
                    raise IngestionError(f"{name}: archive compression ratio exceeds the safety limit.")
            if any("vbaproject" in n.lower() for n in names):
                raise IngestionError(f"{name}: macro-enabled documents are not supported.")
    except (zipfile.BadZipFile, EOFError) as exc:
        raise IngestionError(f"{name}: invalid Office document archive.") from exc


def _parts(text: str, limit: int = 4500) -> list[str]:
    """Split long extracted passages without paraphrasing or losing text."""
    text = text.strip()
    output: list[str] = []
    while len(text) > limit:
        boundary = text.rfind(" ", 0, limit)
        if boundary < limit // 2:
            boundary = limit
        output.append(text[:boundary].strip())
        text = text[boundary:].strip()
    if text:
        output.append(text)
    return output


class _Builder:
    def __init__(self, doc_id: str):
        self.doc_id = doc_id
        self.chunks: list[Chunk] = []
        self.characters = 0

    def add(self, text: str, **metadata: object) -> None:
        text = re.sub(r"\s+", " ", text).strip()
        self.characters += len(text)
        if self.characters > MAX_TEXT_CHARS:
            raise IngestionError("Document exceeds the 500,000 extracted-character limit; split it into smaller documents.")
        for part in _parts(text):
            self.chunks.append(Chunk(
                id=f"{self.doc_id}-c{len(self.chunks) + 1:05d}",
                document_id=self.doc_id, text=part, **metadata,
            ))


def _read_pdf(content: bytes, builder: _Builder) -> list[str]:
    reader = PdfReader(io.BytesIO(content), strict=False)
    if reader.is_encrypted:
        raise IngestionError("Password-protected PDF files are not supported. Upload an unencrypted copy.")
    if len(reader.pages) > MAX_PAGES:
        raise IngestionError(f"PDF exceeds the {MAX_PAGES}-page limit.")
    blank_pages: list[int] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if not text.strip():
            blank_pages.append(page_number)
            continue
        # PDFs often put every word on its own line. Normalize those layout
        # breaks before splitting numbered clauses, including "3.10.Работники".
        text = re.sub(r"\s+", " ", text).strip()
        paragraphs = [p for p in re.split(r"(?<!\S)(?=\d+(?:\.\d+)*[.)](?:\s|[^\W\d_]))", text) if p.strip()]
        for paragraph_number, paragraph in enumerate(paragraphs, start=1):
            clause = re.match(r"(\d+(?:\.\d+)*)[.)]", paragraph)
            builder.add(paragraph.strip(), page=page_number, paragraph=paragraph_number,
                        section=clause.group(1) if clause else None)
    if not builder.chunks:
        raise IngestionError("No readable text was found in this PDF. Scanned/image-only PDFs need OCR before upload.")
    return ([f"No text extracted from PDF pages {', '.join(map(str, blank_pages))}; those pages were not analyzed (OCR may be required)."]
            if blank_pages else [])


def _read_docx(content: bytes, builder: _Builder) -> list[str]:
    document = DocxDocument(io.BytesIO(content))
    section: str | None = None
    paragraph_number = 0
    table_number = 0
    # Iterating the body XML preserves the order of paragraphs and tables.
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph_number += 1
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if paragraph.style and paragraph.style.name.startswith(("Heading", "Title")) and text:
                section = text
            builder.add(text, section=section, paragraph=paragraph_number)
        elif child.tag.endswith("}tbl"):
            table_number += 1
            table = Table(child, document)
            for row_number, row in enumerate(table.rows, start=1):
                text = " | ".join(cell.text.strip() for cell in row.cells)
                if text.strip(" |"):
                    location = f"{section + ' · ' if section else ''}Table {table_number}"
                    builder.add(text, section=location, row=row_number)
    return ["DOCX extraction covers body paragraphs and tables; headers, footers, text boxes, comments, footnotes and tracked changes are not analyzed."]


def _read_xlsx(content: bytes, builder: _Builder) -> list[str]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    row_count = 0
    formulas = 0
    warnings: list[str] = []
    try:
        if len(workbook.worksheets) > 100:
            raise IngestionError("Workbook exceeds the 100-sheet limit.")
        for sheet in workbook.worksheets:
            if sheet.max_row and sheet.max_row > MAX_ROWS:
                raise IngestionError(f"Sheet {sheet.title}: exceeds the {MAX_ROWS:,}-row limit.")
            if sheet.max_column and sheet.max_column > 1000:
                raise IngestionError(f"Sheet {sheet.title}: exceeds the 1,000-column limit.")
            for row_number, row in enumerate(sheet.iter_rows(), start=1):
                row_count += 1
                if row_count > MAX_ROWS:
                    raise IngestionError(f"Workbook exceeds the {MAX_ROWS:,} total-row limit.")
                values: list[str] = []
                for cell in row:
                    if cell.value is None:
                        continue
                    if cell.data_type == "f":
                        formulas += 1
                    values.append(f"{cell.coordinate}: {cell.value}")
                if values:
                    builder.add(" | ".join(values), sheet=sheet.title, row=row_number)
        if formulas:
            warnings.append(f"{formulas} spreadsheet formula(s) retained as text; formulas are never executed or evaluated.")
    finally:
        workbook.close()
    return warnings


def ingest_document(filename: str, content: bytes, side: Literal["before", "after"]) -> Document:
    if side not in ("before", "after"):
        raise IngestionError("Document group must be before or after.")
    name, document_type = validate_file(filename, content)
    doc_id = "doc-" + hashlib.sha256(side.encode() + b"\0" + content).hexdigest()[:24]
    builder = _Builder(doc_id)
    try:
        readers = {"pdf": _read_pdf, "docx": _read_docx, "xlsx": _read_xlsx}
        warnings = readers[document_type](content, builder)
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError(f"{name}: could not read the document. Check that it is valid and unencrypted.") from exc
    if not builder.chunks:
        raise IngestionError(f"{name}: no readable text was found.")
    return Document(id=doc_id, filename=name, document_type=document_type,
                    side=side, chunks=builder.chunks, warnings=warnings)
