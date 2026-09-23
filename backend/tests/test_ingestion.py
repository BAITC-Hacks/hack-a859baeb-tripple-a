import io
import zipfile
from pathlib import Path

import pytest
from app.ingestion import (
    IngestionError,
    ingest_document,
    sanitize_filename,
    validate_file,
)
from docx import Document as DocxDocument
from openpyxl import Workbook
from pypdf import PdfWriter


def docx_bytes():
    doc = DocxDocument()
    doc.add_heading("1. Responsibilities", level=1)
    doc.add_paragraph("The Audit Department reviews internal controls.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Risk Department"
    table.cell(0, 1).text = "Maintains the risk register"
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def test_supported_file_validation():
    assert validate_file("Policy.DOCX", docx_bytes()) == ("Policy.DOCX", "docx")
    for name, contents in (("a.exe", b"anything"), ("a.pdf", b""), ("a.pdf", b"not a PDF"), ("a.docx", b"broken")):
        with pytest.raises(IngestionError):
            validate_file(name, contents)


def test_filename_is_sanitized():
    assert sanitize_filename("../../secrets/contract.docx") == "contract.docx"
    assert sanitize_filename("C:\\temp\\contract.docx") == "contract.docx"
    assert "\n" not in sanitize_filename("injected\nname.pdf")


def test_docx_normalized_structure_and_stable_ids():
    content = docx_bytes()
    document = ingest_document("policy.docx", content, "before")
    assert document.document_type == "docx"
    assert document.side == "before"
    assert len(document.chunks) == 3
    assert document.chunks[1].section == "1. Responsibilities"
    assert document.chunks[1].paragraph == 2
    assert document.chunks[2].section == "1. Responsibilities · Table 1"
    assert document.chunks[2].row == 1
    assert "risk register" in document.chunks[2].text
    assert all(chunk.document_id == document.id and chunk.text for chunk in document.chunks)
    assert len({chunk.id for chunk in document.chunks}) == 3
    assert ingest_document("renamed.docx", content, "before").id == document.id
    assert ingest_document("policy.docx", content, "after").id != document.id


def test_xlsx_preserves_sheet_row_cells_and_does_not_execute_formula():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Responsibilities"
    sheet.append(["Department", "Function"])
    sheet.append(["Audit", "Review controls"])
    sheet.append(["=1+1", None])
    output = io.BytesIO()
    workbook.save(output)
    document = ingest_document("functions.xlsx", output.getvalue(), "after")
    assert document.chunks[1].sheet == "Responsibilities"
    assert document.chunks[1].row == 2
    assert document.chunks[1].text == "A2: Audit | B2: Review controls"
    assert "=1+1" in document.chunks[2].text
    assert any("never executed" in warning for warning in document.warnings)


def test_archive_type_mismatch_and_unsafe_entries():
    with pytest.raises(IngestionError, match="contents do not match"):
        validate_file("renamed.xlsx", docx_bytes())
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", "<xml/>")
        archive.writestr("[Content_Types].xml", "<xml/>")
        archive.writestr("../../escape", "bad")
    with pytest.raises(IngestionError, match="unsafe"):
        validate_file("unsafe.docx", buffer.getvalue())


def test_oversized_and_high_compression_files_are_rejected(monkeypatch):
    monkeypatch.setattr("app.ingestion.MAX_FILE_BYTES", 100)
    with pytest.raises(IngestionError, match="per-file limit"):
        validate_file("oversize.pdf", b"%PDF-" + b"A" * 100)
    monkeypatch.setattr("app.ingestion.MAX_FILE_BYTES", 20 * 1024 * 1024)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "A" * 2_000_000)
        archive.writestr("[Content_Types].xml", "<xml/>")
    with pytest.raises(IngestionError, match="compression ratio"):
        validate_file("high-ratio.docx", buffer.getvalue())


def test_image_only_pdf_is_rejected_with_ocr_guidance():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(IngestionError, match="OCR"):
        ingest_document("scan.pdf", buffer.getvalue(), "before")


def test_real_fixture_clause_and_page_metadata():
    path = Path(__file__).resolve().parents[2] / "fixtures" / "before.pdf"
    if not path.exists():
        pytest.skip("Private audit fixtures are optional")
    document = ingest_document(path.name, path.read_bytes(), "before")
    clause = next(chunk for chunk in document.chunks if chunk.section == "3.4" and chunk.page == 6)
    assert "Департамент непрерывного мониторинга" in clause.text
    assert "Департамент контроля качества" in clause.text
    assert "\n" not in clause.text
    assert any(chunk.section == "3.10" for chunk in document.chunks)
