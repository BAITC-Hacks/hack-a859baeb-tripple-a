import io
from types import SimpleNamespace

import pytest
from app import ai, main
from app.evidence import validate_evidence
from app.local import local_findings, local_functions, local_units
from app.models import Chunk, Document, Evidence, Finding, FunctionMapping, UnitMapping
from app.pipeline import (
    run_pipeline,
    semantic_review,
    validate_findings,
    validate_mappings,
)
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from pydantic import ValidationError


def documents():
    result = []
    for side in ("before", "after"):
        content = "Audit Department reviews access controls."
        result.append(
            Document(
                id=side,
                filename=f"{side}.pdf",
                document_type="pdf",
                side=side,
                chunks=[
                    Chunk(
                        id=f"{side}-c1",
                        document_id=side,
                        text=content,
                        page=1,
                        section="1.1",
                    )
                ],
            )
        )
    return result


def ref(side="before", quote="Audit Department"):
    return Evidence(chunk_id=f"{side}-c1", quote=quote)


def issue(**overrides):
    args = {
        "id": "issue",
        "category": "loss",
        "title": "Potential missing review",
        "explanation": "Requires review",
        "severity": "medium",
        "confidence": 0.5,
        "evidence": [ref()],
        "recommendation": "Confirm the AFTER owner.",
    }
    return Finding(**{**args, **overrides})


def test_evidence_rejects_fabricated_quote_chunk_wrong_document_and_side():
    docs = documents()
    assert validate_evidence([ref()], docs, {"before"})
    assert not validate_evidence([ref(quote="invented responsibility")], docs)
    assert not validate_evidence([Evidence(chunk_id="missing", quote="Audit")], docs)
    assert not validate_evidence([ref()], docs, {"after"})
    assert not validate_evidence([], docs)
    docs[0].chunks[0].document_id = "wrong"
    assert not validate_evidence([ref()], docs)


def test_schema_rejects_invalid_confidence_category_extra_fields():
    with pytest.raises(ValidationError):
        issue(confidence=1.01)
    with pytest.raises(ValidationError):
        issue(category="confirmed_crime")
    with pytest.raises(ValidationError):
        ai.UnitExtraction.model_validate({"units": [], "secret": "unexpected"})


def test_findings_are_rejected_or_downgraded_and_always_reviewable():
    docs, warnings = documents(), []
    invalid = issue(evidence=[ref(quote="forged")])
    duplicate = issue(
        category="duplication", evidence=[ref("after")], requires_review=False
    )
    accepted = validate_findings([invalid, duplicate], docs, warnings)
    assert len(accepted) == 1
    assert accepted[0].category == "uncertain"
    assert accepted[0].requires_review
    assert "1 findings were rejected" in warnings[0]


def test_mapping_rejects_wrong_entity_ids_and_backfills_coverage():
    docs = documents()
    units = local_units(docs)
    mapping = UnitMapping(
        id="bad",
        before_ids=["forged"],
        after_ids=[],
        status="removed",
        confidence=0.99,
        explanation="unsupported",
        evidence=[ref()],
    )
    warnings = []
    output = validate_mappings([mapping], units, docs, "unit", warnings)
    assert len(output) == 2
    assert all(m.status == "uncertain" for m in output)
    assert warnings


def test_number_changes_do_not_mean_lost_functions():
    docs = documents()
    docs[0].chunks[0].text = "1.2. Audit Department reviews access controls."
    docs[1].chunks[0].text = "8.9. Audit Department reviews access controls."
    result = run_pipeline(docs, "local", lambda *args: None)
    assert result.summary["units_before"] == 1
    assert result.summary["functions_before"] == 1
    assert result.function_mappings[0].match_type == "unchanged"
    assert result.summary["lost_functions"] == 0
    assert "Local lexical review" in result.report
    assert "[E1]" in result.report
    assert "page 1, section 1.1" in result.report


def test_conflict_policy_is_not_an_actual_conflict():
    docs = documents()
    for d in docs:
        d.chunks[0].text = "Audit Department must prevent conflicts of interest."
    functions = local_functions(docs, local_units(docs))
    assert not any(f.category == "conflict" for f in local_findings(functions, []))


def test_independent_reviewer_rejects_claim_and_limits_mapping_confidence(monkeypatch):
    docs = documents()
    units = local_units(docs)
    mapping = UnitMapping(
        id="mapping",
        before_ids=[units[0].id],
        after_ids=[units[1].id],
        status="unchanged",
        confidence=0.99,
        explanation="Same name",
        evidence=[ref(), ref("after")],
    )
    verdicts = ai.EvidenceReview(
        verdicts=[
            ai.Verdict(
                finding_id="issue",
                verdict="unsupported",
                reason="Evidence does not establish absence.",
                confidence=0.9,
            ),
            ai.Verdict(
                finding_id="mapping",
                verdict="supported",
                reason="Same name only.",
                confidence=0.6,
            ),
        ]
    )
    monkeypatch.setattr(ai, "ask", lambda *args: verdicts)
    warnings = []
    assert semantic_review([issue()], [mapping], [], docs, warnings) == []
    assert mapping.confidence == 0.6
    assert "rejected" in warnings[0]


def test_real_local_api_vertical_slice(tmp_path, monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)

    def make_doc(text):
        doc = DocxDocument()
        doc.add_paragraph(text)
        out = io.BytesIO()
        doc.save(out)
        return out.getvalue()

    with TestClient(main.create_app(tmp_path / "real.sqlite3")) as client:
        response = client.post(
            "/api/analyses",
            data={"title": "Integration test", "mode": "local"},
            files=[
                (
                    "before_files",
                    (
                        "before.docx",
                        make_doc("Audit Department reviews access controls."),
                    ),
                ),
                (
                    "after_files",
                    (
                        "after.docx",
                        make_doc("Audit Department reviews access controls."),
                    ),
                ),
            ],
        )
        assert response.status_code == 202
        result = client.get("/api/analyses/" + response.json()["id"]).json()
        assert result["status"] == "completed", result
        assert result["result"]["summary"]["functions_before"] == 1
        assert result["result"]["summary"]["units_after"] == 1
        assert result["result"]["function_mappings"][0]["match_type"] == "unchanged"


def test_openai_client_uses_structured_output_and_does_not_store(monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "test-placeholder")
    captured = {}

    def parse(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(output_parsed=ai.UnitExtraction(units=[]))

    monkeypatch.setattr(
        ai,
        "OpenAI",
        lambda **kwargs: SimpleNamespace(responses=SimpleNamespace(parse=parse)),
    )
    assert ai.ask(ai.UnitExtraction, "Extract units", {"text": "example"}).units == []
    assert captured["store"] is False
    assert captured["text_format"] is ai.UnitExtraction
    assert "UNTRUSTED DATA" in captured["input"][0]["content"]


def test_staged_openai_pipeline_with_mocked_structured_responses(monkeypatch):
    docs = documents()
    calls = []

    def fake(schema, instruction, payload):
        calls.append(schema.__name__)
        if schema is ai.UnitExtraction:
            side = payload["side"]
            return ai.UnitExtraction(
                units=[ai.ExtractedUnit(name="Audit Department", evidence=[ref(side)])]
            )
        if schema is ai.FunctionExtraction:
            side = payload["side"]
            return ai.FunctionExtraction(
                functions=[
                    ai.ExtractedFunction(
                        unit_id=payload["units"][0]["id"],
                        description=docs[0].chunks[0].text,
                        evidence=[ref(side, docs[0].chunks[0].text)],
                    )
                ]
            )
        if schema is ai.UnitMatches:
            return ai.UnitMatches(
                mappings=[
                    UnitMapping(
                        id="m",
                        before_ids=[payload[0]["id"]],
                        after_ids=[payload[1]["id"]],
                        status="unchanged",
                        confidence=0.95,
                        explanation="Same named unit.",
                        evidence=[ref(), ref("after")],
                    )
                ]
            )
        if schema is ai.FunctionMatches:
            return ai.FunctionMatches(
                mappings=[
                    FunctionMapping(
                        id="f",
                        before_ids=[payload["before"][0]["id"]],
                        after_ids=[payload["after"][0]["id"]],
                        match_type="unchanged",
                        confidence=0.95,
                        explanation="Same source duty.",
                        evidence=[ref(), ref("after")],
                    )
                ]
            )
        if schema is ai.IssueDetection:
            return ai.IssueDetection(
                findings=[issue(evidence=[ref(quote="fabricated source")])]
            )
        if schema is ai.EvidenceReview:
            return ai.EvidenceReview(
                verdicts=[
                    ai.Verdict(
                        finding_id=item["claim"]["id"],
                        verdict="supported",
                        reason="Matches sources.",
                        confidence=0.9,
                    )
                    for item in payload
                ]
            )
        raise AssertionError(schema)

    monkeypatch.setattr(ai, "ask", fake)
    stages = []
    result = run_pipeline(docs, "openai", lambda stage, label: stages.append(stage))
    assert stages == list(range(1, 8))
    assert result.summary["units_before"] == 1
    assert result.summary["functions_after"] == 1
    assert not result.findings
    assert any("rejected" in w for w in result.warnings)
    assert {
        "UnitExtraction",
        "FunctionExtraction",
        "UnitMatches",
        "FunctionMatches",
        "IssueDetection",
        "EvidenceReview",
    } <= set(calls)


def test_local_owner_heading_handles_inflection_and_page_footer():
    doc = Document(
        id="d",
        filename="before.pdf",
        document_type="pdf",
        side="before",
        chunks=[
            Chunk(
                id="c1",
                document_id="d",
                text="Департамент непрерывного мониторинга (ДНМ).",
                page=1,
            ),
            Chunk(
                id="c2",
                document_id="d",
                text="5.4. Директор департамента непрерывного мониторинга: 8",
                page=8,
            ),
            Chunk(
                id="c3",
                document_id="d",
                text="5.4.1. организует непрерывный мониторинг качества деятельности.",
                page=9,
            ),
        ],
    )
    units = local_units([doc])
    functions = local_functions([doc], units)
    assert len(functions) == 1
    assert functions[0].unit_id == units[0].id
    assert any(e.chunk_id == "c2" for e in functions[0].evidence)


def test_shared_after_candidate_is_uncertain_not_a_loss():
    from app.local import local_function_matches
    from app.models import Function

    text = "Review access controls every quarter."
    funcs = [
        Function(
            id="b1",
            unit_id="u1",
            side="before",
            description=text,
            evidence=[ref(quote=text)],
        ),
        Function(
            id="b2",
            unit_id="u2",
            side="before",
            description=text,
            evidence=[ref(quote=text)],
        ),
        Function(
            id="a1",
            unit_id="u3",
            side="after",
            description=text,
            evidence=[ref("after", text)],
        ),
    ]
    mappings = local_function_matches(funcs)
    assert len(mappings) == 2
    assert {m.match_type for m in mappings} == {"unchanged", "uncertain"}
    assert all(m.match_type != "potential_loss" for m in mappings)
