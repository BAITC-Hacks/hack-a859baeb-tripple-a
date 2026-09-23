"""Extract immutable fragments with format-specific source locations."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile

MAX_TEXT = 250_000

def extract(name: str, raw: bytes):
    ext = Path(name).suffix.lower()
    if ext not in {'.txt', '.docx', '.pdf', '.xlsx'}:
        raise ValueError('Поддерживаются .docx, .pdf, .xlsx, .txt. Старые .doc/.xls конвертируйте.')
    if ext in {'.docx', '.xlsx'}:
        try:
            with ZipFile(BytesIO(raw)) as z:
                if sum(x.file_size for x in z.infolist()) > 60_000_000:
                    raise ValueError('Распакованный файл превышает 60 МБ')
        except BadZipFile as e:
            raise ValueError('Повреждённый Office-файл') from e
    fragments, warnings = [], []
    total = 0
    def add(text, locator):
        nonlocal total
        text = text.strip()
        if not text: return
        total += len(text)
        if total > MAX_TEXT: raise ValueError('Документ превышает лимит 250 000 символов')
        fragments.append({'text': text, 'locator': locator})
    if ext == '.txt':
        for i, line in enumerate(raw.decode('utf-8-sig').splitlines(), 1):
            add(line, {'line': i})
    elif ext == '.docx':
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc = Document(BytesIO(raw))
        for i, block in enumerate(doc.iter_inner_content(), 1):
            if isinstance(block, Paragraph): add(block.text, {'block': i, 'kind': 'paragraph'})
            elif isinstance(block, Table):
                for row_i, row in enumerate(block.rows, 1):
                    add(' | '.join(c.text for c in row.cells), {'block': i, 'row': row_i, 'kind': 'table'})
        warnings.append('DOCX: координаты блоков/строк; страницы, колонтитулы и текстовые поля не извлекаются.')
    elif ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(raw))
        if reader.is_encrypted: raise ValueError('Зашифрованный PDF не поддерживается')
        if len(reader.pages) > 300: raise ValueError('Не более 300 страниц PDF')
        for page_i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ''
            if not text.strip(): warnings.append(f'Страница {page_i}: нет текста, требуется OCR.')
            for line_i, line in enumerate(text.splitlines(), 1):
                add(line, {'page': page_i, 'line': line_i})
    else:
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(raw), read_only=True, data_only=False)
        for sheet in wb:
            if (sheet.max_row or 0) * (sheet.max_column or 0) > 100_000:
                wb.close()
                raise ValueError('XLSX: область листа превышает 100 000 ячеек')
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        add(str(cell.value), {'sheet': sheet.title, 'cell': cell.coordinate})
        wb.close()
        warnings.append('XLSX: читаются значения и формулы ячеек; изображения и схемы не распознаются.')
    if not fragments: raise ValueError('Текст не найден. Для сканов предварительно выполните OCR.')
    return fragments, warnings
