import os
import tempfile
from pathlib import Path
from io import BytesIO
import copy
import pytest
from fastapi.testclient import TestClient
from pypdf import PdfWriter

TEST_ROOT = tempfile.TemporaryDirectory()
os.environ['DATABASE_URL'] = 'sqlite:///' + TEST_ROOT.name + '/test.db'
os.environ['STORAGE_DIR'] = TEST_ROOT.name + '/uploads'
from app.main import app
from app.analysis import demo, validate_and_enrich
from app.parsers import extract

EXAMPLES = Path(__file__).parents[1] / 'examples'

@pytest.fixture
def client():
    with TestClient(app) as c: yield c

def project(c):
    r = c.post('/projects', json={'name':'Test'}); assert r.status_code==201
    return r.json()['id']

def upload(c, pid, phase, content=None):
    raw = content if content is not None else (EXAMPLES / f'{phase}.txt').read_bytes()
    return c.post(f'/projects/{pid}/documents', data={'phase':phase}, files={'file':(phase+'.txt', raw)})

def source_rows():
    rows=[]
    for phase in ['before','after']:
        for i, line in enumerate((EXAMPLES / f'{phase}.txt').read_text().splitlines()):
            rows.append({'id': f'{phase}-{i}', 'document_id':phase, 'document_name':phase+'.txt', 'phase':phase, 'text':line,'locator':{'line':i+1}})
    return rows

def test_end_to_end(client):
    pid=project(client)
    docs = [upload(client,pid,p).json() for p in ['before','after']]
    r=client.post(f'/projects/{pid}/analyses', json={'mode':'demo'})
    assert r.status_code==201, r.text
    body=r.json(); result=body['result']
    kinds=[x['kind'] for x in result['findings']]
    assert kinds.count('potential_loss')==1
    assert kinds.count('potential_duplication')==1
    assert kinds.count('potential_conflict')==1
    assert {x['status'] for x in result['changes']}=={'reorganized','preserved'}
    assert len(result['matches'])==4
    loss=next(x for x in result['findings'] if x['kind']=='potential_loss')
    assert 'реестр поставщиков' in loss['evidence'][0]['quote']
    for finding in result['findings']:
        for e in finding['evidence']:
            assert e['quote'] in result['sources'][e['fragment_id']]['text']
    assert client.get('/analyses/'+body['id']).json()['result']==result
    assert len(client.get(f'/projects/{pid}/analyses').json())==1
    assert client.get(docs[0]['download_url']).content==(EXAMPLES/'before.txt').read_bytes()
    assert client.get('/documents/'+docs[0]['id']+'/fragments').json()[0]['locator']=={'line':1}

def test_upload_deduplicates(client):
    pid=project(client)
    assert upload(client,pid,'before').json()['id']==upload(client,pid,'before').json()['id']
    assert len(client.get(f'/projects/{pid}/documents').json())==1

def test_missing_phase(client):
    pid=project(client); upload(client,pid,'before')
    assert client.post(f'/projects/{pid}/analyses',json={'mode':'demo'}).status_code==422

def test_isolation(client):
    a,b=project(client),project(client)
    upload(client,a,'before'); upload(client,b,'after')
    assert client.post(f'/projects/{a}/analyses',json={'mode':'demo'}).status_code==422

def test_invalid_uploads(client):
    pid=project(client)
    assert upload(client,pid,'bad',b'test').status_code==422
    assert upload(client,pid,'before',b'').status_code==422
    r=client.post(f'/projects/{pid}/documents',data={'phase':'before'},files={'file':('bad.exe',b'bad')})
    assert r.status_code==422
    assert client.get('/analyses/missing').status_code==404

def test_rejects_fabricated_quote():
    rows=source_rows(); report=demo(rows)
    report.functions[0].evidence[0].quote='Invented source text'
    with pytest.raises(ValueError,match='цитату'): validate_and_enrich(report,rows)

def test_rejects_other_project_source():
    rows=source_rows(); report=demo(rows)
    report.units[0].evidence[0].fragment_id='foreign-project-fragment'
    with pytest.raises(ValueError): validate_and_enrich(report,rows)

def test_rejects_incomplete_coverage():
    rows=source_rows(); report=demo(rows); report.matches.pop()
    with pytest.raises(ValueError,match='покрывает'): validate_and_enrich(report,rows)

def test_rejects_wrong_phase():
    rows=source_rows(); report=demo(rows)
    report.matches[0].after_ids=[report.matches[0].before_id]
    with pytest.raises(ValueError,match='фаза'): validate_and_enrich(report,rows)

def test_preserved_report_has_no_findings():
    rows=[r for r in source_rows() if r['phase']=='before']
    rows += [dict(r,id='copy-'+r['id'],phase='after',document_id='copy') for r in list(rows)]
    report=validate_and_enrich(demo(rows),rows)
    assert report['findings']==[]

def test_docx_order_and_table():
    from docx import Document
    d=Document(); d.add_paragraph('Подразделение: Отдел')
    t=d.add_table(rows=1,cols=1); t.cell(0,0).text='1. Проверяет заявки.'
    d.add_paragraph('2. Составляет отчёт.')
    b=BytesIO(); d.save(b)
    fragments,warnings=extract('a.docx',b.getvalue())
    assert len(fragments)==3 and fragments[1]['locator']['kind']=='table'
    assert fragments[2]['text'].startswith('2.')

def test_xlsx_locations():
    from openpyxl import Workbook
    w=Workbook(); w.active.title='Функции'; w.active['B2']='Проверяет заявки'
    b=BytesIO(); w.save(b)
    fragments,_=extract('a.xlsx',b.getvalue())
    assert fragments[0]['locator']=={'sheet':'Функции','cell':'B2'}

def test_pdf_scan_rejected():
    w=PdfWriter(); w.add_blank_page(width=100,height=100); b=BytesIO(); w.write(b)
    with pytest.raises(ValueError,match='OCR'): extract('scan.pdf',b.getvalue())

def test_pdf_text_and_location():
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    w=PdfWriter(); page=w.add_blank_page(width=400,height=400)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):w._add_object(font)})})
    stream=DecodedStreamObject(); stream.set_data(b'BT /F1 12 Tf 10 350 Td (Department finance) Tj ET')
    page[NameObject('/Contents')]=w._add_object(stream)
    b=BytesIO(); w.write(b)
    fragments,_=extract('text.pdf',b.getvalue())
    assert fragments[0]['text']=='Department finance'
    assert fragments[0]['locator']['page']==1

def test_model_integration_contract(client,monkeypatch):
    import app.main as main
    pid=project(client)
    for p in ['before','after']: upload(client,pid,p)
    monkeypatch.setattr(main,'ollama',lambda rows:demo(rows))
    r=client.post(f'/projects/{pid}/analyses',json={'mode':'ollama'})
    assert r.status_code==201 and r.json()['mode']=='ollama'

def test_model_unavailable(client,monkeypatch):
    import app.main as main
    import httpx
    def fail(rows): raise httpx.ConnectError('unavailable')
    monkeypatch.setattr(main,'ollama',fail)
    pid=project(client)
    for p in ['before','after']: upload(client,pid,p)
    assert client.post(f'/projects/{pid}/analyses',json={'mode':'ollama'}).status_code==502

def test_ollama_http_adapter(monkeypatch):
    import httpx
    import json
    import app.analysis as analysis
    rows=source_rows()
    expected=demo(rows)
    real_client=httpx.Client
    def handler(request):
        payload=json.loads(request.content)
        assert request.url.path=='/api/chat'
        assert payload['stream'] is False
        assert payload['format']['type']=='object'
        return httpx.Response(200,json={'message':{'content':expected.model_dump_json()}})
    monkeypatch.setenv('OLLAMA_MODEL','test-local')
    monkeypatch.setattr(analysis.httpx,'Client',lambda **kw:real_client(transport=httpx.MockTransport(handler),**kw))
    assert analysis.ollama(rows)==expected

def test_invalid_model_output_not_saved(client,monkeypatch):
    import app.main as main
    def invalid(rows):
        report=demo(rows)
        report.functions[0].evidence[0].fragment_id='fabricated'
        return report
    monkeypatch.setattr(main,'ollama',invalid)
    pid=project(client)
    for p in ['before','after']: upload(client,pid,p)
    assert client.post(f'/projects/{pid}/analyses',json={'mode':'ollama'}).status_code==422
    assert client.get(f'/projects/{pid}/analyses').json()==[]
